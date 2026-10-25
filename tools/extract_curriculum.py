"""「초중고 세부 목차와 학습 목표」 docx에서 학년별 학습 항목과 331개 단원 연결표를 추출한다.

사용: python3 -I tools/extract_curriculum.py <목차.docx> curriculum/
"""
import json
import re
import sys
from pathlib import Path

import docx

GRADES = [
    ("G01", "초1", "초등학교 1학년", "E12"), ("G02", "초2", "초등학교 2학년", "E12"),
    ("G03", "초3", "초등학교 3학년", "E34"), ("G04", "초4", "초등학교 4학년", "E34"),
    ("G05", "초5", "초등학교 5학년", "E56"), ("G06", "초6", "초등학교 6학년", "E56"),
    ("G07", "중1", "중학교 1학년", "M"), ("G08", "중2", "중학교 2학년", "M"),
    ("G09", "중3", "중학교 3학년", "M"), ("G10", "고1", "고등학교 1학년", "H"),
    ("G11", "고2", "고등학교 2학년", "H"), ("G12", "고3", "고등학교 3학년", "H"),
]
# 13개 계통: 학년표의 계통명, 연결표의 절 번호, 코드
STRANDS = [
    ("KOR", "국어", "국어와 의사소통", 1), ("MAT", "수학", "수학", 2), ("ENG", "영어", "영어", 3),
    ("SOC", "사회·역사", "사회와 역사", 4), ("ETH", "도덕·윤리", "도덕과 윤리", 5),
    ("SCI", "과학", "자연과학", 6), ("TEC", "기술·가정", "기술과 공학 및 가정생활", 7),
    ("INF", "정보·AI", "정보와 컴퓨터 및 인공지능", 8), ("ART", "음악·미술", "음악과 미술", 9),
    ("PHE", "체육·보건·안전", "체육과 보건 및 안전", 10),
    ("HUM", "철학·논리·심리·교육·종교", "철학과 논리 및 심리와 교육과 종교", 11),
    ("CAR", "진로·직업", "진로와 직업", 12), ("ENV", "환경·생태", "환경과 생태 및 지속가능성", 13),
]
STRAND_BY_NAME = {s[1]: s for s in STRANDS}
GRADE_BY_SHORT = {g[1]: g for g in GRADES}


def heading_level(p):
    name = p.style.name if p.style is not None else ""
    m = re.match(r"Heading (\d)", name)
    return int(m.group(1)) if m else None


def iter_blocks(document):
    """본문의 문단과 표를 문서 순서대로 돌려준다."""
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", docx.text.paragraph.Paragraph(child, document)
        elif tag == "tbl":
            yield "t", docx.table.Table(child, document)


def parse_grade_range(text):
    """'초1–초2, 중1–중3' 같은 표기를 학년 코드 목록으로 바꾼다."""
    order = [g[1] for g in GRADES]
    result = []
    for part in re.split(r"[,，]\s*", text.strip()):
        part = part.strip()
        if not part:
            continue
        ends = re.split(r"\s*[–—-]+\s*", part)
        if len(ends) == 1:
            span = [ends[0]]
        else:
            a, b = order.index(ends[0]), order.index(ends[1])
            span = order[a:b + 1]
        for short in span:
            code = GRADE_BY_SHORT[short][0]
            if code not in result:
                result.append(code)
    return result


def main(src, out_dir):
    document = docx.Document(src)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    grade_goals = {}
    items = []
    units = []
    current_grade = None
    current_unit_strand = None
    in_units = False

    for kind, block in iter_blocks(document):
        if kind == "p":
            text = block.text.strip()
            level = heading_level(block)
            if level == 1:
                in_units = text.startswith("기존 331개 단원")
                current_grade = None
                for g in GRADES:
                    if text.startswith(g[2]):
                        current_grade = g
                if text.startswith("자료 제작을 위한") or text.startswith("근거 자료"):
                    in_units = False
            elif level == 2 and in_units:
                m = re.match(r"(\d+)\s+(.+)", text)
                current_unit_strand = next((s for s in STRANDS if s[3] == int(m.group(1))), None) if m else None
            elif current_grade and text.startswith("학년 목표"):
                grade_goals[current_grade[0]] = text.replace("학년 목표", "", 1).strip()
            continue

        table = block
        header = [c.text.strip() for c in table.rows[0].cells]
        if current_grade and header[:2] == ["계통과 형태", "학습 내용과 목표"]:
            for row in table.rows[1:]:
                left, right = row.cells[0], row.cells[1]
                strand_name = left.paragraphs[0].text.strip()
                form = " ".join(p.text.strip() for p in left.paragraphs[1:] if p.text.strip())
                strand = STRAND_BY_NAME[strand_name]
                for p in right.paragraphs:
                    if not p.text.strip():
                        continue
                    title = "".join(r.text for r in p.runs if r.bold).strip()
                    goal = "".join(r.text for r in p.runs if not r.bold).strip()
                    if title == "정규 교과 없음":
                        continue  # 초1·2 영어처럼 학습 항목이 없는 자리 표시 행
                    seq = sum(1 for it in items if it["grade"] == current_grade[0] and it["strand"] == strand[0]) + 1
                    items.append({
                        "id": f"KR22-{current_grade[0]}-{strand[0]}-I{seq:02d}",
                        "grade": current_grade[0],
                        "grade_label": current_grade[1],
                        "band": current_grade[3],
                        "strand": strand[0],
                        "strand_label": strand[1],
                        "form": form,
                        "title": title,
                        "goal": goal,
                    })
        elif in_units and current_unit_strand and header[0].startswith("기존 번호"):
            for row in table.rows[1:]:
                name_cell, grade_cell = row.cells[0].text.strip(), row.cells[1].text.strip()
                m = re.match(r"(\d+\.\d+\.\d+)\s+(.+)", name_cell)
                if not m:
                    continue
                units.append({
                    "unit_id": m.group(1),
                    "strand": current_unit_strand[0],
                    "field_id": m.group(1).rsplit(".", 1)[0],
                    "name": m.group(2),
                    "grades_text": grade_cell,
                    "grades": parse_grade_range(grade_cell),
                })

    meta = {
        "curriculum": "2022 개정 교육과정 계열 / 권장 학년 배치 v1",
        "source": "초중고 세부 목차와 학습 목표 (기준일 2026-10-14)",
        "note": "학년 배치와 목표 문장은 교육자료 제작용 권장안이며 공식 성취기준 원문이나 번호 대조표가 아니다.",
        "grades": [{"code": g[0], "label": g[1], "name": g[2], "band": g[3], "goal": grade_goals.get(g[0], "")} for g in GRADES],
        "strands": [{"code": s[0], "label": s[1], "name": s[2], "unit_section": s[3]} for s in STRANDS],
    }
    (out / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / "grade_items.json").write_text(json.dumps(items, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / "units.json").write_text(json.dumps(units, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = {u["field_id"] for u in units}
    print(f"grades={len(grade_goals)} items={len(items)} units={len(units)} fields={len(fields)} "
          f"grade×strand={len({(i['grade'], i['strand']) for i in items})}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
